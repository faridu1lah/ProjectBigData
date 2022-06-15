import streamlit as st
import base64
from pptx import Presentation
from io import BytesIO
import streamlit.components.v1 as components


def load_view():
    
    
    st.title("Download center for the developers!")
    # with st.form(key="my_form"):
      # Opening file from file path
    with open('data/business_case.pdf', "rb") as f:
        base64_pdf = base64.b64encode(f.read()).decode('utf-8')
    col11, col22 = st.columns([3,1])
    with col11:
        # Embedding PDF in HTML
        pdf_display = F'<iframe src="data:application/pdf;base64,{base64_pdf}" title="hallo" width="900" height="600" type="application/pdf">'

        # Displaying File
        st.markdown(pdf_display, unsafe_allow_html=True)

        prs = Presentation()
        # title_slide_layout = prs.slide_layouts[0]
        # slide = prs.slides.add_slide(title_slide_layout)
        # title = slide.shapes.title
        # subtitle = slide.placeholders[1]

        # title.text = "Hello, World!"
        # subtitle.text = "python-pptx was here!"

    # save the output into binary form
        binary_output = BytesIO()
        prs.save(binary_output) 

        # st.download_button(label='Download ppw',
        #                     data=binary_output.getvalue(),
        #                     file_name='ADS.pptx') 
    with col22:
        st.header("Download file's ")
        st.download_button(label='Sprint 1',data=binary_output.getvalue(),file_name='data/Sprint1.pptx') 
        st.download_button(label='Sprint 2',data=binary_output.getvalue(),file_name='data/Sprint2.pptx') 
        st.download_button(label='Sprint 4',data=binary_output.getvalue(),file_name='data/Sprint4.pptx') 
        st.download_button(label='Business case',data=binary_output.getvalue(),file_name='data/business_case.pdf') 
        st.download_button(label='Scrum process',data=binary_output.getvalue(),file_name='data/scrum_pr.pdf') 
        st.write("[master.zip](https://github.com/FaridullahKhorshid/ProjectBigData/archive/refs/heads/master.zip)")
   